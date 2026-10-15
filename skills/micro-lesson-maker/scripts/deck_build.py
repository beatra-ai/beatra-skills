#!/usr/bin/env python3
"""Write an editable 16:9 .pptx from a lesson outline, draw the same outline as
one PNG per slide, and download the images a finished image task produced.
Writing needs python-pptx, drawing needs Pillow (python-pptx installs it), and
downloading uses only the standard library. This script never starts another
program.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import urllib.error
import urllib.request
from pathlib import Path

LAYOUTS = ("title", "bullets", "image")
DOWNLOAD_LIMIT = 64 * 1024 * 1024
# python-pptx's default template: "Title Slide" and "Title Only".
TITLE_SLIDE, TITLE_ONLY = 0, 5
BULLET = "\u2022"

# Where each part of a slide sits, in inches on the 13.333 x 7.5 slide; the
# .pptx and the rendered PNG use the same boxes.
TITLE_BOX = (1.0, 2.6, 11.333, 1.4)
SUBTITLE_BOX = (1.0, 4.1, 11.333, 0.8)
HEADING_BOX = (0.6, 0.4, 12.1, 1.0)
BODY_TOP, BODY_HEIGHT = 1.6, 5.4
SIDE_PICTURE_BOX = (7.5, 1.6, 5.2, 5.4)
FULL_PICTURE_BOX = (0.6, 1.6, 12.1, 5.5)
SIZES = {"title": 48, "subtitle": 24, "heading": 36, "body": 24}
PIXELS_PER_INCH = 144  # 13.333 in -> 1920 px
FONT_FILES = {
    "darwin": [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ],
    "win32": [
        "%WINDIR%\\Fonts\\msyh.ttc",
        "%WINDIR%\\Fonts\\simhei.ttf",
        "%WINDIR%\\Fonts\\Deng.ttf",
        "%WINDIR%\\Fonts\\simsun.ttc",
        "%WINDIR%\\Fonts\\arial.ttf",
    ],
    "linux": [
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ],
}
FONT_FOLDERS = {
    "darwin": ["/System/Library/Fonts", "/Library/Fonts", "~/Library/Fonts", "/System/Library/AssetsV2"],
    "win32": ["%WINDIR%\\Fonts", "%LOCALAPPDATA%\\Microsoft\\Windows\\Fonts"],
    "linux": ["/usr/share/fonts", "/usr/local/share/fonts", "~/.local/share/fonts", "~/.fonts"],
}
CJK_FONT_NAMES = ("cjk", "wqy", "yahei", "pingfang", "hiragino")
FONT_SUFFIXES = (".ttf", ".ttc", ".otf", ".otc")
FONT_SEARCH_DEPTH, FONT_SEARCH_FILES = 4, 20000
FONT_HELP = "pass --font <path> with a .ttf, .otf, or .ttc font that has these characters"


class OutlineError(Exception):
    """An outline or image task the Agent should fix before building."""


def _read_json(path: Path, what: str) -> dict:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise OutlineError(f"{what} {path.name} cannot be read as JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise OutlineError(f"{what} {path.name} must be a JSON object")
    return data


def load_outline(path: Path) -> dict:
    outline = _read_json(path, "the outline")
    slides = outline.get("slides")
    if not isinstance(slides, list) or not slides:
        raise OutlineError("the outline needs a non-empty slides list")
    base = path.parent
    background = outline.get("background")
    if background and not (base / background).is_file():
        raise OutlineError(f"background {background} not found next to the outline")
    for number, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            raise OutlineError(f"slide {number}: must be an object")
        if slide.get("layout") not in LAYOUTS:
            raise OutlineError(f"slide {number}: layout must be one of {', '.join(LAYOUTS)}")
        if not str(slide.get("title", "")).strip():
            raise OutlineError(f"slide {number}: the title is empty")
        image = slide.get("image")
        if slide["layout"] == "image" and not image:
            raise OutlineError(f"slide {number}: an image slide needs an image")
        if image and not (base / image).is_file():
            raise OutlineError(f"slide {number}: {image} not found next to the outline")
    return outline


def _inches(box):
    from pptx.util import Inches

    return [Inches(value) for value in box]


def _fill(frame, text, size, bold=False):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    frame.word_wrap = True
    frame.text = text
    for paragraph in frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold


def _place(shape, box) -> None:
    shape.left, shape.top, shape.width, shape.height = _inches(box)


def _bullets(slide, items, box, size):
    """A text box whose paragraphs are real bullets, so PowerPoint and WPS indent and restyle them."""

    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    left, top, width, height = _inches(box)
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    for number, item in enumerate(items):
        paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.space_after = Pt(size // 2)
        for run in paragraph.runs:
            run.font.size = Pt(size)
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(Inches(0.4)))
        properties.set("indent", str(-Inches(0.4)))
        properties.append(properties.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        properties.append(properties.makeelement(qn("a:buChar"), {"char": BULLET}))


def _picture(slide, path, box):
    left, top, width, height = _inches(box)
    picture = slide.shapes.add_picture(str(path), left, top, height=height)
    if picture.width > width:
        ratio = width / picture.width
        picture.width = int(picture.width * ratio)
        picture.height = int(picture.height * ratio)
    picture.left = int(left + (width - picture.width) / 2)
    picture.top = int(top + (height - picture.height) / 2)


def _to_back(slide, shape) -> None:
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    tree.insert(2, shape._element)  # after the group's own nvGrpSpPr and grpSpPr


def _body_box(spec: dict) -> tuple:
    width = 6.6 if spec.get("image") else 12.1
    return (0.6, BODY_TOP, width, BODY_HEIGHT)


def build(outline_path: Path, out: Path) -> int:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise OutlineError("writing .pptx needs python-pptx: python3 -m pip install --user python-pptx") from exc
    outline = load_outline(outline_path)
    base = outline_path.parent
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for spec in outline["slides"]:
        layout = TITLE_SLIDE if spec["layout"] == "title" else TITLE_ONLY
        slide = deck.slides.add_slide(deck.slide_layouts[layout])
        title = slide.shapes.title
        if spec["layout"] == "title":
            _place(title, TITLE_BOX)
            _fill(title.text_frame, spec["title"], SIZES["title"], bold=True)
            subtitle = next(shape for shape in slide.placeholders if shape.placeholder_format.idx == 1)
            if spec.get("subtitle"):
                _place(subtitle, SUBTITLE_BOX)
                _fill(subtitle.text_frame, spec["subtitle"], SIZES["subtitle"])
            else:
                subtitle._element.getparent().remove(subtitle._element)
        else:
            _place(title, HEADING_BOX)
            _fill(title.text_frame, spec["title"], SIZES["heading"], bold=True)
            if spec["layout"] == "bullets":
                _bullets(slide, spec.get("bullets", []), _body_box(spec), SIZES["body"])
                if spec.get("image"):
                    _picture(slide, base / spec["image"], SIDE_PICTURE_BOX)
            else:
                _picture(slide, base / spec["image"], FULL_PICTURE_BOX)
        if outline.get("background"):
            background = slide.shapes.add_picture(
                str(base / outline["background"]), 0, 0, deck.slide_width, deck.slide_height
            )
            _to_back(slide, background)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    out.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(out))
    return len(outline["slides"])


def _pillow():
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError as exc:
        raise OutlineError(
            "drawing slide images needs Pillow, which python-pptx installs: "
            "python3 -m pip install --user python-pptx"
        ) from exc
    return Image, ImageDraw, ImageFont


def _expand(entry: str) -> Path:
    return Path(os.path.expanduser(os.path.expandvars(entry)))


def _platform_key(platform: str) -> str:
    if platform.startswith("win"):
        return "win32"
    return "darwin" if platform == "darwin" else "linux"


def font_candidates(platform: str | None = None):
    """Installed fonts worth trying for slide text, best first."""

    key = _platform_key(platform or sys.platform)
    seen: set[str] = set()
    for entry in FONT_FILES.get(key, []):
        path = _expand(entry)
        if path.is_file() and str(path) not in seen:
            seen.add(str(path))
            yield path
    checked = 0
    for folder in (_expand(entry) for entry in FONT_FOLDERS.get(key, [])):
        if not folder.is_dir():
            continue
        for current, children, files in os.walk(folder):
            depth = len(Path(current).relative_to(folder).parts)
            children[:] = sorted(children) if depth < FONT_SEARCH_DEPTH else []
            for name in sorted(files):
                checked += 1
                if checked > FONT_SEARCH_FILES:
                    return
                lowered = name.lower()
                path = Path(current) / name
                if lowered.endswith(FONT_SUFFIXES) and any(token in lowered for token in CJK_FONT_NAMES):
                    if str(path) not in seen:
                        seen.add(str(path))
                        yield path


def _open_font(path: Path, size: int):
    _, _, image_font = _pillow()
    face = 0
    if path.suffix.lower() in (".ttc", ".otc"):
        for index in range(16):
            try:
                family = image_font.truetype(str(path), size, index=index).getname()[0] or ""
            except OSError:
                break
            if not family.startswith(".") and family.split()[-1:] == ["SC"]:
                face = index
                break
    try:
        return image_font.truetype(str(path), size, index=face)
    except OSError as exc:
        raise OutlineError(f"the font {path} cannot be opened ({exc}); {FONT_HELP}") from exc


def missing_characters(font, text: str) -> str:
    """The characters of text, once each, that the font draws as its missing-glyph shape."""

    image, draw, _ = _pillow()

    def shape(char: str) -> tuple:
        canvas = image.new("L", (96, 96))
        draw.Draw(canvas).text((0, 0), char, font=font, fill=255)
        return font.getlength(char), canvas.tobytes()

    absent = shape("\U0010fffd")  # a private-use code point no font maps
    return "".join(
        char for char in dict.fromkeys(text) if char.isprintable() and not char.isspace() and shape(char) == absent
    )


def _outline_text(outline: dict) -> str:
    parts = []
    for spec in outline["slides"]:
        parts += [str(spec.get("title", "")), str(spec.get("subtitle", "")), *map(str, spec.get("bullets", []))]
    return "".join(parts) + BULLET


def choose_font(outline: dict, font: Path | None, candidates=None) -> Path:
    """--font as given, else the first installed font that draws every character of the outline."""

    if font is not None:
        if not font.is_file():
            raise OutlineError(f"--font {font} is not a font file; {FONT_HELP}")
        return font
    text = _outline_text(outline)
    best, best_gap = None, None
    for candidate in candidates if candidates is not None else font_candidates():
        try:
            gap = missing_characters(_open_font(candidate, 48), text)
        except OutlineError:
            continue
        if not gap:
            return candidate
        if best_gap is None or len(gap) < len(best_gap):
            best, best_gap = candidate, gap
    if best is None:
        raise OutlineError(f"no font for the slide text was found on this computer; {FONT_HELP}")
    raise OutlineError(f"no installed font draws every character of the slides ({best} lacks {best_gap}); {FONT_HELP}")


def _px(value: float) -> int:
    return int(round(value * PIXELS_PER_INCH))


def _wrap(draw, text: str, font, width: int) -> list[str]:
    """Lines that fit width: Latin words stay whole, CJK breaks between characters."""

    lines: list[str] = []
    line = ""
    for token in _tokens(text):
        trial = line + token
        if line and draw.textlength(trial.rstrip(), font=font) > width:
            lines.append(line.rstrip())
            line = token.lstrip()
        else:
            line = trial
    if line.strip():
        lines.append(line.rstrip())
    return lines or [""]


def _tokens(text: str) -> list[str]:
    tokens: list[str] = []
    word = ""
    for char in text:
        if char.isascii() and not char.isspace():
            word += char
            continue
        if word:
            tokens.append(word)
            word = ""
        tokens.append(char)
    if word:
        tokens.append(word)
    return tokens


def _draw_text(draw, text, font, box, bold=False, indent=0) -> None:
    left, top, width, _ = (_px(value) for value in box)
    line_height = int(font.size * 1.25)
    y = top
    for line in _wrap(draw, text, font, width - indent):
        draw.text((left + indent, y), line, font=font, fill="black", stroke_width=1 if bold else 0, stroke_fill="black")
        y += line_height


def _draw_bullets(draw, items, font, box) -> None:
    left, top, width, _ = (_px(value) for value in box)
    indent, line_height = _px(0.4), int(font.size * 1.25)
    y = top
    for item in items:
        draw.text((left, y), BULLET, font=font, fill="black")
        for line in _wrap(draw, item, font, width - indent):
            draw.text((left + indent, y), line, font=font, fill="black")
            y += line_height
        y += font.size // 2


def _paste(canvas, path: Path, box) -> None:
    image, _, _ = _pillow()
    left, top, width, height = (_px(value) for value in box)
    with image.open(path) as source:
        picture = source.convert("RGBA")
    scale = min(width / picture.width, height / picture.height)
    picture = picture.resize((max(1, int(picture.width * scale)), max(1, int(picture.height * scale))))
    canvas.paste(picture, (left + (width - picture.width) // 2, top + (height - picture.height) // 2), picture)


def render(outline_path: Path, out_dir: Path, font: Path | None = None) -> list[str]:
    """Draw every slide of the outline as page-NNN.png, 1920x1080, in the boxes the .pptx uses."""

    outline = load_outline(outline_path)
    if out_dir.exists() and any(out_dir.iterdir()):
        raise OutlineError(f"{out_dir} is not empty; render into an empty pages folder")
    image, draw_module, _ = _pillow()
    font_path = choose_font(outline, font)
    fonts = {name: _open_font(font_path, size * 2) for name, size in SIZES.items()}  # 1 pt = 2 px here
    base = outline_path.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for number, spec in enumerate(outline["slides"], start=1):
        canvas = image.new("RGBA", (_px(13.333), _px(7.5)), "white")
        if outline.get("background"):
            with image.open(base / outline["background"]) as source:
                canvas.paste(source.convert("RGBA").resize(canvas.size), (0, 0))
        draw = draw_module.Draw(canvas)
        if spec["layout"] == "title":
            _draw_text(draw, spec["title"], fonts["title"], TITLE_BOX, bold=True)
            if spec.get("subtitle"):
                _draw_text(draw, spec["subtitle"], fonts["subtitle"], SUBTITLE_BOX)
        else:
            _draw_text(draw, spec["title"], fonts["heading"], HEADING_BOX, bold=True)
            if spec["layout"] == "bullets":
                _draw_bullets(draw, spec.get("bullets", []), fonts["body"], _body_box(spec))
                if spec.get("image"):
                    _paste(canvas, base / spec["image"], SIDE_PICTURE_BOX)
            else:
                _paste(canvas, base / spec["image"], FULL_PICTURE_BOX)
        target = out_dir / f"page-{number:03d}.png"
        canvas.convert("RGB").save(target, format="PNG")
        written.append(target.name)
    return written


def image_urls(result: dict) -> list[str]:
    data = result.get("structuredContent") or result
    task = data.get("task") or data
    status = data.get("status") or task.get("status")
    if status != "succeeded":
        raise OutlineError(f"task {data.get('task_id') or task.get('task_id')} is {status}; nothing to download")
    output = task.get("output") or data.get("output") or {}
    candidates = list(output.get("images") or []) + list(data.get("artifacts") or task.get("artifacts") or [])
    urls = []
    for item in candidates:
        url = item.get("url") if isinstance(item, dict) else item
        if isinstance(url, str) and url.startswith("https://") and url not in urls:
            urls.append(url)
    if not urls:
        raise OutlineError("the finished task returned no image URL")
    return urls


def image_suffix(head: bytes) -> str:
    if head.startswith(b"\x89PNG"):
        return ".png"
    if head.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return ".webp"
    raise OutlineError("the downloaded file is not a PNG, JPEG, or WebP image")


class _HttpsOnly(urllib.request.HTTPRedirectHandler):
    """Refuse a redirect to anything but https before it is followed."""

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        if not newurl.startswith("https://"):
            raise OutlineError("the result URL redirected away from https")
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def _opener():
    return urllib.request.build_opener(_HttpsOnly())


def fetch(url: str, folder: Path, stem: str) -> Path:
    """Download one image as <stem>.png or <stem>.jpg; WebP becomes PNG."""

    if not url.startswith("https://"):
        raise OutlineError("only https result URLs are downloaded")
    request = urllib.request.Request(url, headers={"User-Agent": "beatra-deck"})
    try:
        with _opener().open(request, timeout=120) as response:
            if not response.geturl().startswith("https://"):
                raise OutlineError("the result URL redirected away from https")
            data = response.read(DOWNLOAD_LIMIT + 1)
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        raise OutlineError(f"the image could not be downloaded: {exc}") from exc
    if len(data) > DOWNLOAD_LIMIT:
        raise OutlineError("the download is larger than an image should be")
    folder.mkdir(parents=True, exist_ok=True)
    suffix = image_suffix(data[:16])
    if suffix != ".webp":
        target = folder / f"{stem}{suffix}"
        target.write_bytes(data)
        return target
    try:
        from PIL import Image  # installed with python-pptx
    except ImportError as exc:
        raise OutlineError("a WebP image needs Pillow, which python-pptx installs") from exc
    import io

    target = folder / f"{stem}.png"
    Image.open(io.BytesIO(data)).save(target, format="PNG")
    return target


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="deck_build.py")
    commands = parser.add_subparsers(dest="command", required=True)
    make = commands.add_parser("build", help="write an editable .pptx from an outline")
    make.add_argument("outline", type=Path)
    make.add_argument("--out", type=Path, required=True)
    draw = commands.add_parser("render", help="draw one PNG per slide of an outline, without PowerPoint")
    draw.add_argument("outline", type=Path)
    draw.add_argument("--out-dir", type=Path, required=True)
    draw.add_argument("--font", type=Path)
    images = commands.add_parser("fetch-images", help="download the images of a finished image task")
    images.add_argument("task_result", type=Path)
    images.add_argument("--dir", type=Path, required=True)
    images.add_argument("--name", required=True)
    args = parser.parse_args(argv)
    try:
        if args.command == "build":
            count = build(args.outline, args.out)
            print(json.dumps({"slides": count, "out": str(args.out)}, ensure_ascii=False))
        elif args.command == "render":
            pages = render(args.outline, args.out_dir, args.font)
            print(json.dumps({"pages": len(pages), "out_dir": str(args.out_dir)}, ensure_ascii=False))
        else:
            urls = image_urls(_read_json(args.task_result, "the task result"))
            saved = [str(fetch(url, args.dir, f"{args.name}-{number}")) for number, url in enumerate(urls, start=1)]
            print(json.dumps({"images": saved}, ensure_ascii=False))
    except (OutlineError, OSError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
